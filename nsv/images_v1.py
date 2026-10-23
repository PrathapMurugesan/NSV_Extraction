"""
images.py - D4 defect images -> one table per image group (cracking, potholes) with lane, chainage, lat/long.

    d4/Crack Images_DHMS/L1/cb7c960f-..._lcms_345690.jpg
    d4/Pothole Images/L2/73480728-..._lcms_345700_1.jpg
       ^ group folder    ^ lane ^ survey run id  ^ LCMS distance (m) - NOT the chainage

The number after '_lcms_' is the distance counter of the LCMS run, offset from the road chainage by a
constant per run (e.g. L1 +200 m, R1 +150 m). The offset is learned from the NSV Cracks workbook, which
lists 'Image Name' next to its real 'Start Chainage':
    chainage = number in the file name + offset of its run id

Where the location of an image comes from (first one that works):
    1. GPS in the image EXIF (if the vendor geotagged it)
          -> nearest Defect Report row of that lane (by its start lat/long) gives the chainage
    2. chainage (as above)
          -> the Defect Report row of that chainage gives the lane's lat/long
             (interpolated between the row's start and end point)

The Defect Report 'Data Sheet' has, per 10 m row: Chainage Start / End, then per lane (L4..R4)
Start Latitude, Start Longitude, End Latitude, End Longitude.
"""
import io
import re
from pathlib import PurePosixPath

import numpy as np
import openpyxl
import pandas as pd

import settings
from nsv.log import get_logger

logger = get_logger("nsv.images")

IMAGE_EXTENSIONS = (".jpg", ".jpeg", ".png", ".tif", ".tiff")
LANE_LIMIT_M = 50              # lanes of one defect-report row further apart than this = suspicious coordinates
IMAGE_NAME = re.compile(r"^(.+?)_lcms_(\d+(?:\.\d+)?)", re.IGNORECASE)     # <run id>_lcms_<distance>
COLUMNS = ["image_name", "lane", "image_number", "chainage_offset", "chainage", "chainage_start", "chainage_end",
           "latitude", "longitude", "location_source", "gps_distance_m", "image_path"]


# =============================================================================
# the one function d4 calls
# =============================================================================
def map_images(run, defect_report_path, groups, cracks_path=None, sheet_name="Data Sheet"):
    """
    groups = {"cracking": ["crack"], "potholes": ["pothole"]}   group key -> keywords of its folder name
    cracks_path = NSV Cracks workbook (Image Name + Start Chainage) -> chainage offset per survey run
    Saves d4_images_<group>.parquet + .csv per group that has images; summary in run.manifest["images"].
    """
    lanes = read_lane_locations(settings.INPUT.read_bytes(defect_report_path), sheet_name)
    logger.info(f"[images] defect report: {len(lanes)} lane(s) with locations {sorted(lanes)}")

    offsets = read_offsets(settings.INPUT.read_bytes(cracks_path)) if cracks_path else {}
    logger.info(f"[images] chainage offset per survey run (from NSV Cracks): {offsets}")

    found = _image_files(run.input_folder, groups)
    run.manifest["images"] = {"chainage_offsets": offsets}
    for group, files in found.items():
        rows = [_locate(path, lane, lanes, offsets) for path, lane in files]
        df = pd.DataFrame(rows, columns=COLUMNS)
        unmapped = int(df["latitude"].isna().sum())
        no_offset = int(df["chainage_offset"].isna().sum())
        saved = run.save_table(df, "images", group) if rows else []
        run.manifest["images"][group] = {"images": len(rows), "unmapped": unmapped,
                                         "without_offset": no_offset, "saved": saved}
        logger.info(f"[images] {group}: {len(rows)} image(s), {unmapped} without a location")
        if unmapped:
            logger.warning(f"[images] {group}: {unmapped} image(s) could not be mapped - see location_source")
        if no_offset:
            logger.warning(f"[images] {group}: {no_offset} image(s) of a run not in NSV Cracks - the file-name "
                           f"number was used as chainage, location may be off")


# =============================================================================
# NSV Cracks -> {run id: offset in m}   (Start Chainage - number in the image name)
# =============================================================================
def read_offsets(data):
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    wb.close()
    head = [_clean(v) for v in rows[0]]
    name_col, chainage_col = head.index("image name"), head.index("start chainage")

    found = {}
    for r in rows[1:]:
        match = IMAGE_NAME.match(str(r[name_col] or ""))
        start = _number(r[chainage_col])
        if match and start is not None:
            found.setdefault(match.group(1), []).append(start - float(match.group(2)))
    # most common difference per run (the first frame of a run can differ by a few metres)
    return {run_id: float(pd.Series(diffs).mode()[0]) for run_id, diffs in found.items()}


# =============================================================================
# Defect Report -> {lane: DataFrame(start, end, start_lat, start_lon, end_lat, end_lon)}
# =============================================================================
def read_lane_locations(data, sheet_name):
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    top, sub = [_clean(v) for v in rows[0]], [_clean(v) for v in rows[1]]

    chainage = top.index("chainage")                                # 'Chainage' over 'Start' | 'End'
    lane_cols = {top[i].upper(): i for i in range(len(top))         # 'L1' over 'Start Latitude' ...
                 if top[i] and sub[i] == "start latitude"}

    lanes = {}
    for lane, col in lane_cols.items():
        table = []
        for r in rows[2:]:
            start, end = _number(r[chainage]), _number(r[chainage + 1])
            lat1, lon1, lat2, lon2 = (_number(r[col + k]) for k in range(4))
            if start is not None and lat1 is not None and lon1 is not None:
                table.append((start, end, lat1, lon1, lat2, lon2))
        if table:
            lanes[lane] = pd.DataFrame(table, columns=["start", "end", "lat1", "lon1", "lat2", "lon2"])
    return _fix_reversed_lanes(lanes)


def _fix_reversed_lanes(lanes, limit_m=LANE_LIMIT_M):
    """
    Lanes of one row lie a few metres apart. A lane whose points are far from the other lanes of the same
    row but close once its rows are reversed was delivered in reverse chainage order (seen: R2 in the
    Defect Report) -> its points are reversed here, with a warning. Only for the image mapping; the
    defect report table itself is saved as delivered.
    """
    for lane, table in lanes.items():
        others = [t.set_index("start")[["lat1", "lon1"]] for name, t in lanes.items() if name != lane]
        if not others:
            continue
        ref = pd.concat(others).groupby(level=0).median()               # middle of the other lanes, per row
        here = table.set_index("start")[["lat1", "lon1"]]
        flipped = table[["lat2", "lon2"]].iloc[::-1].set_axis(table["start"]).set_axis(["lat1", "lon1"], axis=1)
        apart = lambda pts: _distance_m(ref["lat1"], ref["lon1"], pts["lat1"], pts["lon1"]).dropna().median()  # noqa: E731
        now, reversed_ = apart(here), apart(flipped)
        if now <= limit_m:
            continue
        if reversed_ <= limit_m:
            rev = table.iloc[::-1]
            table[["lat1", "lon1", "lat2", "lon2"]] = rev[["lat2", "lon2", "lat1", "lon1"]].to_numpy()
            logger.warning(f"[images] defect report lane {lane}: lat/long are in REVERSE chainage order "
                           f"({now:.0f} m from the other lanes, {reversed_:.0f} m once reversed) - reversed for "
                           f"the image mapping")
        else:
            logger.warning(f"[images] defect report lane {lane}: lat/long are {now:.0f} m from the other lanes "
                           f"of the same row - its image locations may be wrong")
    return lanes


# =============================================================================
# one image -> one row
# =============================================================================
def _locate(path, lane, lanes, offsets):
    name = PurePosixPath(path).name
    match = IMAGE_NAME.match(name)
    number = float(match.group(2)) if match else None
    offset = offsets.get(match.group(1)) if match else None
    chainage = number + (offset or 0) if number is not None else None
    row = {"image_name": name, "lane": lane, "image_number": number, "chainage_offset": offset,
           "chainage": chainage, "image_path": path}
    table = lanes.get(lane)
    if table is None:
        return {**row, "location_source": f"lane '{lane}' has no locations in the defect report"}

    # ---- 1. GPS in the image ----
    gps = _exif_gps(settings.INPUT.read_bytes(path))
    if gps:
        distance = _distance_m(gps[0], gps[1], table["lat1"], table["lon1"])
        best = table.loc[distance.idxmin()]
        return {**row, "chainage_start": best["start"], "chainage_end": best["end"],
                "latitude": gps[0], "longitude": gps[1], "location_source": "exif_gps",
                "gps_distance_m": round(float(distance.min()), 2)}

    # ---- 2. chainage -> defect report row ----
    if chainage is None:
        return {**row, "location_source": "no GPS in image and no chainage in file name"}
    inside = table[(table["start"] <= chainage) & (chainage < table["end"])]     # 345890 -> row 345890-345900
    if inside.empty:
        inside = table[table["end"] == chainage]                                # the very last chainage
    if inside.empty:
        return {**row, "location_source": f"chainage outside the defect report "
                                          f"({table['start'].min():g} - {table['end'].max():g})"}
    best = inside.iloc[0]
    lat, lon = best["lat1"], best["lon1"]
    length = best["end"] - best["start"] if pd.notna(best["end"]) else 0
    if length > 0 and pd.notna(best["lat2"]) and pd.notna(best["lon2"]):
        f = (chainage - best["start"]) / length                      # position inside the 10 m row
        lat, lon = lat + (best["lat2"] - lat) * f, lon + (best["lon2"] - lon) * f
    source = "defect_report_by_chainage" if offset is not None else "defect_report_by_file_number_no_offset"
    return {**row, "chainage_start": best["start"], "chainage_end": best["end"],
            "latitude": lat, "longitude": lon, "location_source": source}


# =============================================================================
# helpers
# =============================================================================
def _image_files(folder, groups):
    """{group: [(path, lane), ...]} - group = first sub-folder whose name has a keyword, lane = folder below it."""
    found = {group: [] for group in groups}
    for path in settings.INPUT.list_files(folder):
        parts = PurePosixPath(path).relative_to(folder.strip("/")).parts
        if len(parts) < 2 or PurePosixPath(path).suffix.lower() not in IMAGE_EXTENSIONS:
            continue
        group = next((g for g, words in groups.items() if any(w in parts[0].lower() for w in words)), None)
        if group:
            lane = parts[1].upper() if len(parts) > 2 else None
            found[group].append((path, lane))
    return found


def _exif_gps(data):
    """(lat, lon) from the image EXIF, or None (also None when Pillow is not installed)."""
    try:
        from PIL import Image
        gps = Image.open(io.BytesIO(data)).getexif().get_ifd(0x8825)
        if not gps or 2 not in gps or 4 not in gps:
            return None
        to_deg = lambda d: float(d[0]) + float(d[1]) / 60 + float(d[2]) / 3600   # noqa: E731
        lat, lon = to_deg(gps[2]), to_deg(gps[4])
        return (-lat if gps.get(1) == "S" else lat, -lon if gps.get(3) == "W" else lon)
    except Exception:
        return None


def _distance_m(lat, lon, lats, lons):
    """Approximate metres (fine at lane scale) between points; each argument a number or a column."""
    dy = (lats - lat) * 111_320
    dx = (lons - lon) * 111_320 * np.cos(np.radians(lat))
    return (dx ** 2 + dy ** 2) ** 0.5


def _clean(value):
    return str(value).strip().lower() if value is not None else ""


def _number(value):
    try:
        return float(value) if value is not None and str(value).strip() != "" else None
    except ValueError:
        return None
