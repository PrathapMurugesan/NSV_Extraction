"""
excel.py - read Excel workbooks. Reading only: no checks, no saving.

    wb = ExcelWorkbook(file_name, data_bytes)
    wb.sheet_names
    table = wb.read_table(sheet_name, options)     -> Table (header found, merged cells handled)
    wb.read_cells(sheet_name, columns="A:G")       -> every non-empty cell, for report-style sheets
    wb.read_formulas(sheet_name, table)            -> {column: {excel_row: "=formula"}}

Sheet options (set per sheet in deliverables/dN.py):
    header_row   Excel row where the header starts (1 = first row). Not set = found automatically.
    header_rows  how many rows the header has (default 1). Several rows are combined into one name
                 per column: 'Current Cycle | Roughness BI (in mm/km) | L1'.
    info_rows    rows inside the header that are legends / notes, not names (1 = first header row).
                 Their text goes to data_dictionary.json instead of the column name.
    skip_rows    rows right under the header to drop (e.g. a row of notes).
    columns      only read these Excel columns, e.g. "A:G" (ignores side tables / dropdown lists).

Merged cells: a merged cell's text belongs to EVERY cell it covers - left to right and top to
bottom. Only real merges are filled; a blank cell that is not merged stays blank. (The old code
filled every blank from the left, which produced names like 'Lane R4 | Remark | Longitude'.)
"""
import io
import re
import warnings
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass, field

import pandas as pd
from openpyxl.utils import get_column_letter, range_boundaries

import settings
from nsv.log import get_logger

logger = get_logger("nsv.excel")

# openpyxl warns about Excel features it can't read (dropdown lists etc.); they don't affect the data
warnings.filterwarnings("ignore", message=".*extension is not supported.*", module="openpyxl")


# =============================================================================
# Small helpers
# =============================================================================
def is_blank(value):
    """None, NaN and empty / whitespace-only text count as blank."""
    if value is None:
        return True
    if isinstance(value, float) and value != value:
        return True
    return isinstance(value, str) and value.strip() == ""


def clean_text(value):
    """Line breaks and repeated spaces -> one space, trimmed. ('NH\\nNumber ' -> 'NH Number')"""
    return re.sub(r"\s+", " ", str(value)).strip()


def column_range(text, width):
    """'A:G' -> (0, 6)  zero-based, inclusive. None -> all columns."""
    if not text:
        return 0, width - 1
    first, last = text.split(":")
    c1 = range_boundaries(f"{first}1:{first}1")[0] - 1
    c2 = range_boundaries(f"{last}1:{last}1")[0] - 1
    return c1, min(c2, width - 1)


# =============================================================================
# The result of reading one sheet as a table
# =============================================================================
@dataclass
class Table:
    columns: list                  # [{"name": "Chainage | Start", "letter": "B", "index": 1}, ...]
    rows: list                     # data values, one list per row, same order as columns
    excel_rows: list               # Excel row number of every data row (-> source_row)
    header_excel_rows: list        # Excel rows used for the header, e.g. [1, 2, 3, 4]
    dictionary: dict = field(default_factory=dict)   # {column name: {"header row 2": "2500"}}

    @property
    def names(self):
        return [c["name"] for c in self.columns]

    def rename(self, names):
        """Give the columns their final names; the original header stays in col['vendor_header']."""
        old_to_new = {}
        for col, name in zip(self.columns, names):
            col.setdefault("vendor_header", col["name"])
            old_to_new[col["name"]] = name
            col["name"] = name
        self.dictionary = {old_to_new.get(k, k): v for k, v in self.dictionary.items()}

    def to_frame(self):
        """DataFrame with 'source_row' first, values exactly as read (not converted)."""
        df = pd.DataFrame(self.rows, columns=self.names, dtype=object)
        df.insert(0, "source_row", self.excel_rows)
        return df


# =============================================================================
# One workbook
# =============================================================================
class ExcelWorkbook:
    def __init__(self, file_name, data):
        self.file_name = file_name
        self.data = data
        self._excel = pd.ExcelFile(io.BytesIO(data))           # values (cached results of formulas)
        self.sheet_names = list(self._excel.sheet_names)
        self._zip = _open_zip(data)                              # None for old .xls files
        self._xml_paths = _sheet_xml_paths(self._zip) if self._zip else {}
        self._formula_book = None                                # opened only if a sheet has formulas

    def close(self):
        self._excel.close()
        if self._formula_book is not None:
            self._formula_book.close()

    # -------------------------------------------------------------------------
    # raw cell values: grid[0] = Excel row 1, grid[r][0] = column A
    # -------------------------------------------------------------------------
    def read_grid(self, sheet):
        df = pd.read_excel(self._excel, sheet_name=sheet, header=None, dtype=object, keep_default_na=False)
        return [[None if is_blank(v) else v for v in row] for row in df.values.tolist()]

    # -------------------------------------------------------------------------
    # merged ranges and formula presence, read straight from the sheet's XML (fast, no full load)
    # -------------------------------------------------------------------------
    def _sheet_xml(self, sheet):
        path = self._xml_paths.get(sheet)
        return self._zip.read(path) if path else b""

    def merged_ranges(self, sheet):
        """[(first_row, first_col, last_row, last_col), ...]  1-based, like Excel."""
        refs = re.findall(rb'<(?:\w+:)?mergeCell\s+ref="([A-Z]+\d+(?::[A-Z]+\d+)?)"', self._sheet_xml(sheet))
        result = []
        for ref in refs:
            c1, r1, c2, r2 = range_boundaries(ref.decode())
            result.append((r1, c1, r2, c2))
        return result

    def merged_refs(self, sheet):
        return [ref.decode() for ref in re.findall(rb'<(?:\w+:)?mergeCell\s+ref="([^"]+)"', self._sheet_xml(sheet))]

    def has_formulas(self, sheet):
        return re.search(rb"<(?:\w+:)?f[\s>/]", self._sheet_xml(sheet)) is not None

    # -------------------------------------------------------------------------
    # read one sheet as a table
    # -------------------------------------------------------------------------
    def read_table(self, sheet, options=None):
        options = options or {}
        grid = self.read_grid(sheet)
        n_header = options.get("header_rows", 1)
        merges = self.merged_ranges(sheet) if n_header > 1 else []   # only multi-row headers need merges
        return build_table(grid, merges,
                           header_row=options.get("header_row"),
                           header_rows=n_header,
                           info_rows=options.get("info_rows", []),
                           skip_rows=options.get("skip_rows", 0),
                           columns=options.get("columns"))

    # -------------------------------------------------------------------------
    # report-style sheets: every non-empty cell with its address
    # -------------------------------------------------------------------------
    def read_cells(self, sheet, columns=None):
        grid = self.read_grid(sheet)
        width = max((len(r) for r in grid), default=0)
        if width == 0:
            return []
        c1, c2 = column_range(columns, width)
        cells = []
        for r, row in enumerate(grid):
            for c in range(c1, min(c2, len(row) - 1) + 1):
                if not is_blank(row[c]):
                    cells.append({"cell": f"{get_column_letter(c + 1)}{r + 1}", "value": row[c]})
        return cells

    # -------------------------------------------------------------------------
    # formulas of the table's data cells (needs a second, formula-aware read)
    # -------------------------------------------------------------------------
    def read_formulas(self, sheet, table):
        """{column name: {excel_row: '=IF(J5>...)'}} - only cells that hold a formula."""
        if not table.excel_rows or not self.has_formulas(sheet):
            return {}
        if self._formula_book is None:
            import openpyxl
            self._formula_book = openpyxl.load_workbook(io.BytesIO(self.data), read_only=True, data_only=False)
        ws = self._formula_book[sheet]
        wanted = set(table.excel_rows)
        first, last = min(wanted), max(wanted)
        result = {c["name"]: {} for c in table.columns}
        for offset, row in enumerate(ws.iter_rows(min_row=first, max_row=last, values_only=True)):
            excel_row = first + offset
            if excel_row not in wanted:
                continue
            for col in table.columns:
                value = row[col["index"]] if col["index"] < len(row) else None
                text = getattr(value, "text", value)            # array formulas come as objects
                if isinstance(text, str) and text.startswith("="):
                    result[col["name"]][excel_row] = text
        return {name: rows for name, rows in result.items() if rows}


# =============================================================================
# Build the table from the raw grid (pure function - easy to test)
# =============================================================================
def build_table(grid, merges, header_row=None, header_rows=1, info_rows=(), skip_rows=0, columns=None):
    width = max((len(r) for r in grid), default=0)
    if width == 0:
        raise ValueError("sheet is empty")
    grid = [row + [None] * (width - len(row)) for row in grid]           # all rows the same length
    c1, c2 = column_range(columns, width)
    ncols = c2 - c1 + 1

    # ---- 1. where does the header start? ----
    if header_row:
        top = header_row - 1
    else:
        top = _find_header_row([row[c1:c2 + 1] for row in grid], settings.HEADER_SCAN_ROWS)
    if top >= len(grid):
        raise ValueError(f"header row {top + 1} is below the last row of the sheet")
    bottom = min(top + header_rows - 1, len(grid) - 1)

    # ---- 2. copy the header block and fill REAL merged cells (both directions) ----
    block = [list(grid[r][c1:c2 + 1]) for r in range(top, bottom + 1)]
    for (mr1, mc1, mr2, mc2) in merges:
        r0, k0 = mr1 - 1, mc1 - 1                                          # top-left cell, 0-based
        if not (top <= r0 <= bottom):
            continue                                                       # merge starts outside the header
        value = grid[r0][k0]
        for r in range(r0, min(mr2 - 1, bottom) + 1):
            for c in range(max(k0, c1), min(mc2 - 1, c2) + 1):
                block[r - top][c - c1] = value

    # ---- 3. one name per column: header rows top to bottom, info rows left out ----
    info = {i - 1 for i in info_rows}
    names, notes = [], []
    for j in range(ncols):
        parts = []
        for i, row in enumerate(block):
            if i in info or is_blank(row[j]):
                continue
            text = clean_text(row[j])
            if text not in parts:                                          # vertical merges repeat the text
                parts.append(text)
        names.append(" | ".join(parts))
        notes.append({f"header row {i + 1}": clean_text(block[i][j])
                      for i in sorted(info) if i < len(block) and not is_blank(block[i][j])})

    # ---- 4. data rows: everything under the header (minus skip_rows) that is not completely empty ----
    rows, excel_rows = [], []
    for r in range(bottom + 1 + skip_rows, len(grid)):
        values = grid[r][c1:c2 + 1]
        if any(not is_blank(v) for v in values):
            rows.append(values)
            excel_rows.append(r + 1)

    # ---- 5. keep columns that have a name or data; name the nameless ones by their Excel letter ----
    keep = [j for j in range(ncols) if names[j] or any(not is_blank(row[j]) for row in rows)]
    table_columns, seen = [], {}
    for j in keep:
        letter = get_column_letter(c1 + j + 1)
        name = names[j] or f"unnamed_{letter}"
        seen[name] = seen.get(name, 0) + 1
        if seen[name] > 1:
            name = f"{name}_{seen[name]}"                                  # duplicates get _2, _3
        table_columns.append({"name": name, "letter": letter, "index": c1 + j})

    dictionary = {col["name"]: notes[j] for col, j in zip(table_columns, keep) if notes[j]}
    return Table(columns=table_columns,
                 rows=[[row[j] for j in keep] for row in rows],
                 excel_rows=excel_rows,
                 header_excel_rows=list(range(top + 1, bottom + 2)),
                 dictionary=dictionary)


def _find_header_row(rows, scan_rows, min_fill=0.6):
    """
    Automatic header detection (single-row headers):
    the first row that is 'full' (at least 60 % of the widest row near the top) and mostly text.
    A merged title row only has text in its first cell, so it is skipped.
    """
    top = rows[:scan_rows]
    filled = [sum(not is_blank(v) for v in row) for row in top]
    widest = max(filled, default=0)
    if widest == 0:
        raise ValueError("sheet is empty")
    for idx, row in enumerate(top):
        values = [v for v in row if not is_blank(v)]
        if len(values) >= widest * min_fill and sum(isinstance(v, str) for v in values) / len(values) >= 0.8:
            return idx
    raise ValueError(f"no header row found in the first {scan_rows} rows")


# =============================================================================
# xlsx internals: which XML file belongs to which sheet
# =============================================================================
def _open_zip(data):
    try:
        return zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:                       # .xls (old binary format): no merges / formulas info
        return None


def _sheet_xml_paths(zf):
    """{'Comparison as per CA': 'xl/worksheets/sheet3.xml', ...}"""
    try:
        book = ET.fromstring(zf.read("xl/workbook.xml"))
        rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    except KeyError:
        return {}
    targets = {r.get("Id"): r.get("Target") for r in rels}
    paths = {}
    for el in book.iter():
        if el.tag.endswith("}sheet"):
            rel_id = next((v for k, v in el.attrib.items() if k.endswith("}id")), None)
            target = targets.get(rel_id, "")
            paths[el.get("name")] = target.lstrip("/") if target.startswith("/") else f"xl/{target}"
    return paths
