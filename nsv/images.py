"""
images.py - survey images (e.g. d4/crack_images_DHMS/L1/*.jpg): one row per image with every detail found.

Per image, in this order:
  1. file facts     file name, full path, lane (from a folder named L1..L4 / R1..R4), size
  2. EXIF           camera date/time, GPS, make/model - only if the vendor embedded it (instant)
  3. OCR            the text printed on the image (e.g. the blue banner of LCMS images)
                    - only when EXIF has no usable GPS lat/long (settings.IMAGE_OCR = "when_no_gps");
                      "always" reads every image
                    - only the regions in settings.IMAGE_TEXT_REGIONS (default: top banner)
                    - read by LABEL, not position: the text is split into fields where there is a gap
                      between columns, and every 'Label: value' becomes a field. A vendor with another
                      layout or other labels still gets all its fields (columns named after its labels).
                    - known labels (settings.IMAGE_FIELDS) get a standard name: latitude, chainage, ...
  4. checks         lat / long in India, valid date, numbers where numbers belong, file name vs text
                    (survey id), lane folder vs lane in the text. Problems -> 'issues' column.
  5. location       (after OCR, see locate) chainage + lat/long, each from its most reliable source:
                    - chainage  = number in the file name  <run id>_lcms_<number>[_n].jpg  (exact, no OCR)
                    - lat/long  = 1. EXIF GPS, 2. Defect Report row of that lane + chainage, 3. OCR banner
                    The OCR banner values are kept as ocr_latitude / ocr_longitude / ocr_chainage and checked
                    against the chosen location (OCR misreads digits: 79.1111 -> 79.1141 = 325 m off).
                    nsv_cracks_chainage = file number + offset of the run in the NSV Cracks workbook (that
                    workbook counts chainage differently: +140..+200 m), to join images to NSV Cracks rows.
Nothing is guessed: a value that is not found or fails a check stays as read / empty, and the raw OCR text
is kept in 'ocr_text' so anyone can see what was on the image.

Speed: OCR ~0.6 s per image per CPU core; images are processed on all cores in parallel
(settings.IMAGE_WORKERS).
"""
import hashlib
import io
import os
import re
from concurrent.futures import ProcessPoolExecutor, ThreadPoolExecutor
from datetime import datetime

from PIL import ExifTags, Image, ImageOps

from nsv.log import get_logger

logger = get_logger("nsv.images")

LANE_FOLDER = re.compile(r"^(?:lane[ _-]?)?([LR][1-4])$", re.IGNORECASE)
LANE_IN_TEXT = re.compile(r"(?<![A-Za-z0-9])([LR][1-4])(?![A-Za-z0-9])")
IMAGE_NAME = re.compile(r"^(.+?)_lcms_(\d+(?:\.\d+)?)", re.IGNORECASE)     # <run id>_lcms_<distance>
LANE_LIMIT_M = 50              # lanes of one defect-report row further apart than this = suspicious coordinates


# =============================================================================
# 1-2. File facts + EXIF
# =============================================================================
def _exif(img):
    """EXIF fields (date, GPS, camera) or {} when the image has none."""
    exif = img.getexif()
    if not exif:
        return {}
    out = {}
    for tag, value in exif.items():
        name = ExifTags.TAGS.get(tag, str(tag))
        if name in ("Make", "Model", "DateTime", "Software", "ImageDescription"):
            out[f"exif_{name.lower()}"] = str(value).strip()
    try:
        sub = exif.get_ifd(0x8769)                                   # Exif IFD: DateTimeOriginal
        if 36867 in sub:
            out["exif_datetime_original"] = str(sub[36867])
    except Exception:
        pass
    try:
        gps = {ExifTags.GPSTAGS.get(k, k): v for k, v in exif.get_ifd(0x8825).items()}
        def degrees(dms, ref):
            d, m, s = (float(x) for x in dms)
            value = d + m / 60 + s / 3600
            return -value if ref in ("S", "W") else value
        if "GPSLatitude" in gps:
            out["exif_latitude"] = degrees(gps["GPSLatitude"], gps.get("GPSLatitudeRef", "N"))
        if "GPSLongitude" in gps:
            out["exif_longitude"] = degrees(gps["GPSLongitude"], gps.get("GPSLongitudeRef", "E"))
        if "GPSAltitude" in gps:
            out["exif_altitude"] = float(gps["GPSAltitude"])
    except Exception:
        pass
    return out


def _gps_ok(row, field_labels):
    """EXIF lat / long present and inside the allowed range (India, settings.IMAGE_FIELDS)."""
    lat, lon = row.get("exif_latitude"), row.get("exif_longitude")
    if lat is None or lon is None:
        return False
    lat_range = field_labels.get("latitude", {}).get("range", (-90, 90))
    lon_range = field_labels.get("longitude", {}).get("range", (-180, 180))
    return lat_range[0] <= lat <= lat_range[1] and lon_range[0] <= lon <= lon_range[1]


# =============================================================================
# 3. OCR by label
# =============================================================================
def _prepare(region, scale=2):
    """Enlarge, grey, invert (light text on dark banner -> dark on light), black & white."""
    big = region.resize((region.width * scale, region.height * scale), Image.LANCZOS)
    gray = ImageOps.grayscale(big)
    if sum(gray.getdata()) / (gray.width * gray.height) < 128:      # dark background -> invert
        gray = ImageOps.invert(gray)
    return gray.point(lambda p: 255 if p > 110 else 0)


def _text_segments(img):
    """OCR -> lines -> split each line where the gap between words is large (= the next column)."""
    import pytesseract
    d = pytesseract.image_to_data(img, config="--psm 6", output_type=pytesseract.Output.DICT)
    lines = {}
    for i, text in enumerate(d["text"]):
        if text.strip():
            key = (d["block_num"][i], d["par_num"][i], d["line_num"][i])
            lines.setdefault(key, []).append((d["left"][i], d["width"][i], d["height"][i], text))
    segments = []
    for key in sorted(lines):
        words = sorted(lines[key])
        current = [words[0]]
        for prev, word in zip(words, words[1:]):
            if word[0] - (prev[0] + prev[1]) > 1.5 * max(prev[2], word[2]):
                segments.append(" ".join(w[3] for w in current))
                current = []
            current.append(word)
        segments.append(" ".join(w[3] for w in current))
    return segments


def _label_key(label):
    return re.sub(r"[^a-z0-9]", "", label.lower())


def _parse_segments(segments, field_labels):
    """'Latitude: 13.32' -> {'latitude': '13.32'}; unknown labels -> 'label_<name>'; no ':' -> unlabelled."""
    known = {_label_key(l): field for field, cfg in field_labels.items() for l in cfg["labels"]}
    fields, unlabelled = {}, []
    for seg in segments:
        if ":" not in seg:
            unlabelled.append(seg)
            continue
        label, value = seg.split(":", 1)
        key = _label_key(label)
        name = known.get(key) or f"label_{re.sub(r'[^a-z0-9]+', '_', label.lower()).strip('_')}"
        fields.setdefault(name, value.strip())
    return fields, unlabelled


# =============================================================================
# 4. Checks
# =============================================================================
def _number(text):
    try:
        return float(str(text).replace(",", ""))
    except (TypeError, ValueError):
        return None


def _checks(row, field_labels):
    issues = []
    for field, cfg in field_labels.items():
        value = row.get(field)
        if value in (None, ""):
            continue
        kind = cfg.get("type")
        if kind == "number":
            n = _number(value)
            if n is None:
                issues.append(f"{field} not a number: '{value}'")
            elif "range" in cfg and not (cfg["range"][0] <= n <= cfg["range"][1]):
                issues.append(f"{field} {n} outside {cfg['range'][0]}-{cfg['range'][1]}")
        elif kind == "date":
            if not any(_parses(value, f) for f in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d")):
                issues.append(f"{field} not a date: '{value}'")

    # file name vs printed text (chainage: see locate)
    stem = os.path.splitext(row["file_name"])[0]
    survey_id = str(row.get("survey_id") or "").replace("...", "").replace("…", "").strip()
    if len(survey_id) >= 8 and survey_id.lower() not in stem.lower():
        issues.append(f"survey id '{survey_id}' not in file name")

    # lane folder vs lane in the description
    lanes_in_text = LANE_IN_TEXT.findall(str(row.get("description") or ""))
    if row.get("lane") and lanes_in_text and row["lane"] not in lanes_in_text:
        issues.append(f"folder lane {row['lane']} but text says {lanes_in_text}")
    return issues


def _parses(text, fmt):
    try:
        datetime.strptime(str(text).strip(), fmt)
        return True
    except ValueError:
        return False


# =============================================================================
# One image (runs in a worker process)
# =============================================================================
def read_image(job):
    """job = (path, image bytes, lane, folder, settings dict) -> one result row (dict)."""
    path, data, lane, folder, cfg = job
    row = {"image_folder": folder, "lane": lane, "file_name": re.split(r"[\\/]", path)[-1], "image_path": path,
           "size_bytes": len(data), "sha256": hashlib.sha256(data).hexdigest()}
    try:
        img = Image.open(io.BytesIO(data))
        img.load()
    except Exception as e:
        row.update(status="unreadable_file", issues=f"cannot open image: {e}")
        return row
    row.update(width=img.width, height=img.height)
    row.update(_exif(img))

    # OCR only when EXIF has no usable GPS (settings.IMAGE_OCR = "when_no_gps"); "always" = every image
    gps_ok = _gps_ok(row, cfg["fields"])
    skip_ocr = cfg.get("ocr_mode", "when_no_gps") == "when_no_gps" and gps_ok
    if cfg["ocr"] and not skip_ocr:
        import pytesseract
        if cfg["tesseract_cmd"]:
            pytesseract.pytesseract.tesseract_cmd = cfg["tesseract_cmd"]
        rgb = img.convert("RGB")
        segments = []
        for _, top, bottom in cfg["regions"]:
            box = (0, int(top * rgb.height), rgb.width, int(bottom * rgb.height))
            segments += _text_segments(_prepare(rgb.crop(box)))
        fields, unlabelled = _parse_segments(segments, cfg["fields"])
        row.update(fields)
        row["ocr_text"] = " | ".join(segments)
        if unlabelled:
            row["ocr_unlabelled"] = " | ".join(unlabelled)

    issues = _checks(row, cfg["fields"])
    found = [f for f in cfg["fields"] if row.get(f) not in (None, "")]
    row["fields_found"] = len(found)
    row["issues"] = "; ".join(issues)
    if skip_ocr:
        row["status"] = "exif_gps (no OCR needed)"
    elif not cfg["ocr"]:
        row["status"] = "not_read (OCR off)"
    elif not found and not any(k.startswith("exif_") for k in row):
        row["status"] = "no_text_found"
    elif not found:
        row["status"] = "exif_only"                      # no text on the image, but EXIF details
    elif issues or len(found) < len(cfg["fields"]):
        row["status"] = "partial"
    else:
        row["status"] = "ok"
    return row


def read_images(paths, make_job, workers, batch=400, progress=None):
    """
    Read all images on several CPU cores, a batch at a time (so 40,000 images never sit in memory at once).
    make_job(path) -> job for read_image.  progress(done, total) is called after every batch.
    Results come back in the same order as paths.
    """
    rows = []
    pool = ProcessPoolExecutor(max_workers=workers) if workers > 1 and len(paths) >= 4 else None
    try:
        for start in range(0, len(paths), batch):
            with ThreadPoolExecutor(max_workers=32) as downloads:    # file reads (S3) in parallel too
                jobs = list(downloads.map(make_job, paths[start:start + batch]))
            rows += list(pool.map(read_image, jobs, chunksize=8)) if pool else [read_image(j) for j in jobs]
            if progress:
                progress(len(rows), len(paths))
    finally:
        if pool:
            pool.shutdown()
    return rows


# =============================================================================
# 5. Location: chainage from the file name, lat/long from EXIF / Defect Report / OCR
#    Runs in the main process after OCR (needs the workbooks, cheap).
# =============================================================================
def locate(row, lanes, offsets, agree_m):
    """
    Adds chainage, latitude, longitude, location_source, location_check_m, nsv_cracks_chainage to one row.
    The OCR values move to ocr_chainage / ocr_latitude / ocr_longitude (kept as read).
    lanes   = read_lane_locations(...)   {lane: table}  ({} = no Defect Report)
    offsets = read_offsets(...)          {run id: m}    ({} = no NSV Cracks)
    """
    for field in ("chainage", "latitude", "longitude"):
        if field in row:
            row[f"ocr_{field}"] = row.pop(field)
    issues = [row["issues"]] if row.get("issues") else []
    ocr_ch = _number(row.get("ocr_chainage"))
    ocr_lat, ocr_lon = _number(row.get("ocr_latitude")), _number(row.get("ocr_longitude"))
    ocr_ok = ocr_lat is not None and ocr_lon is not None and 6 <= ocr_lat <= 38 and 68 <= ocr_lon <= 98

    # ---- chainage: file name first (exact), OCR banner only when the name has none ----
    match = IMAGE_NAME.match(row["file_name"])
    if match:
        chainage = float(match.group(2))
        if ocr_ch is not None and ocr_ch != chainage:
            issues.append(f"OCR chainage {row['ocr_chainage']} differs from file name ({chainage:g})")
        offset = offsets.get(match.group(1))
        if offset is not None:
            row["nsv_cracks_chainage"] = chainage + offset
    else:
        chainage = ocr_ch

    # ---- lat / long ----
    table = lanes.get(row.get("lane"))
    point = _defect_report_point(table, chainage) if table is not None and chainage is not None else None
    if row.get("exif_latitude") is not None and row.get("exif_longitude") is not None:
        lat, lon, source = row["exif_latitude"], row["exif_longitude"], "exif_gps"
    elif point:
        lat, lon, source = point[0], point[1], "defect_report"
    elif ocr_ok:
        lat, lon, source = ocr_lat, ocr_lon, "ocr_banner"
        if table is None:
            issues.append(f"lane '{row.get('lane')}' has no locations in the defect report - OCR lat/long used")
        elif chainage is not None:
            issues.append(f"chainage {chainage:g} outside the defect report - OCR lat/long used")
    else:
        lat = lon = None
        source = "not_found"
        issues.append("no location: no EXIF GPS, no defect report row, no readable OCR lat/long")

    # ---- second opinion: OCR banner vs chosen location ----
    if lat is not None and ocr_ok and source != "ocr_banner":
        check = round(float(_distance_m(lat, lon, ocr_lat, ocr_lon)), 2)
        row["location_check_m"] = check
        if check > agree_m:
            issues.append(f"OCR lat/long {check:.0f} m from {source} (OCR misread?)")
    if point and source == "exif_gps":
        far = float(_distance_m(lat, lon, point[0], point[1]))
        if far > agree_m:
            issues.append(f"EXIF GPS {far:.0f} m from the defect report at chainage {chainage:g}")

    row.update(chainage=chainage, latitude=lat, longitude=lon, location_source=source,
               chainage_start=point[2] if point else None, chainage_end=point[3] if point else None)
    row["issues"] = "; ".join(issues)
    if row.get("status") == "ok" and issues:                       # a location problem -> not ok any more
        row["status"] = "partial"
    return row


def _defect_report_point(table, chainage):
    """(lat, lon, row start, row end) of the lane at chainage, interpolated inside its 10 m row, or None."""
    inside = table[(table["start"] <= chainage) & (chainage < table["end"])]     # 345890 -> row 345890-345900
    if inside.empty:
        inside = table[table["end"] == chainage]                                # the very last chainage
    if inside.empty:
        return None
    best = inside.iloc[0]
    lat, lon = best["lat1"], best["lon1"]
    length = best["end"] - best["start"] if best["end"] == best["end"] else 0    # NaN-safe
    if length > 0 and best["lat2"] == best["lat2"] and best["lon2"] == best["lon2"]:
        f = (chainage - best["start"]) / length
        lat, lon = lat + (best["lat2"] - lat) * f, lon + (best["lon2"] - lon) * f
    return float(lat), float(lon), float(best["start"]), float(best["end"])


def read_offsets(data):
    """NSV Cracks workbook -> {run id: Start Chainage - number in the image name} (most common per run)."""
    import openpyxl
    import pandas as pd
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    wb.close()
    head = [_clean(v) for v in rows[0]]
    if "image name" not in head or "start chainage" not in head:
        logger.warning("[images] NSV Cracks has no 'Image Name' / 'Start Chainage' column - no nsv_cracks_chainage")
        return {}
    name_col, chainage_col = head.index("image name"), head.index("start chainage")
    found = {}
    for r in rows[1:]:
        match = IMAGE_NAME.match(str(r[name_col] or ""))
        start = _number(r[chainage_col])
        if match and start is not None:
            found.setdefault(match.group(1), []).append(start - float(match.group(2)))
    return {run_id: float(pd.Series(diffs).mode()[0]) for run_id, diffs in found.items()}


def read_lane_locations(data, sheet_name):
    """
    Defect Report -> {lane: DataFrame(start, end, lat1, lon1, lat2, lon2)}.
    Row 1: 'Chainage' over Start | End, and the lane name (L1..R4) over 'Start Latitude', Start Longitude,
    End Latitude, End Longitude. Row 2: those sub-headers. Data from row 3.
    """
    import openpyxl
    import pandas as pd
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    top, sub = [_clean(v) for v in rows[0]], [_clean(v) for v in rows[1]]
    if "chainage" not in top:
        logger.warning("[images] defect report has no 'Chainage' column - images located by EXIF / OCR only")
        return {}
    chainage = top.index("chainage")
    lane_cols = {top[i].upper(): i for i in range(len(top)) if top[i] and sub[i] == "start latitude"}

    lanes = {}
    for lane, col in lane_cols.items():
        table = []
        for r in rows[2:]:
            start, end = _number(r[chainage]), _number(r[chainage + 1])
            lat1, lon1, lat2, lon2 = (_number(r[col + k]) for k in range(4))
            if start is not None and lat1 is not None and lon1 is not None:
                table.append((start, end, lat1, lon1, lat2, lon2))
        if table:
            lanes[lane] = pd.DataFrame(table, columns=["start", "end", "lat1", "lon1", "lat2", "lon2"])
    return _fix_reversed_lanes(lanes)


def _fix_reversed_lanes(lanes, limit_m=LANE_LIMIT_M):
    """
    Lanes of one row lie a few metres apart. A lane whose points are far from the other lanes of the same
    row but close once its rows are reversed was delivered in reverse chainage order (seen: R2 in the
    Defect Report) -> its points are reversed here, with a warning. Only for the image mapping; the
    defect report table itself is saved as delivered.
    """
    import pandas as pd
    for lane, table in lanes.items():
        others = [t.set_index("start")[["lat1", "lon1"]] for name, t in lanes.items() if name != lane]
        if not others:
            continue
        ref = pd.concat(others).groupby(level=0).median()               # middle of the other lanes, per row
        here = table.set_index("start")[["lat1", "lon1"]]
        flipped = table[["lat2", "lon2"]].iloc[::-1].set_axis(table["start"]).set_axis(["lat1", "lon1"], axis=1)
        apart = lambda pts: _distance_m(ref["lat1"], ref["lon1"], pts["lat1"], pts["lon1"]).dropna().median()  # noqa: E731
        now, reversed_ = apart(here), apart(flipped)
        if now <= limit_m:
            continue
        if reversed_ <= limit_m:
            rev = table.iloc[::-1]
            table[["lat1", "lon1", "lat2", "lon2"]] = rev[["lat2", "lon2", "lat1", "lon1"]].to_numpy()
            logger.warning(f"[images] defect report lane {lane}: lat/long are in REVERSE chainage order "
                           f"({now:.0f} m from the other lanes, {reversed_:.0f} m once reversed) - reversed for "
                           f"the image mapping")
        else:
            logger.warning(f"[images] defect report lane {lane}: lat/long are {now:.0f} m from the other lanes "
                           f"of the same row - its image locations may be wrong")
    return lanes


def _distance_m(lat, lon, lats, lons):
    """Approximate metres (fine at lane scale) between points; each argument a number or a column."""
    import numpy as np
    dy = (lats - lat) * 111_320
    dx = (lons - lon) * 111_320 * np.cos(np.radians(lat))
    return (dx ** 2 + dy ** 2) ** 0.5


def _clean(value):
    return str(value).strip().lower() if value is not None else ""


def lane_from_path(parts):
    """First folder named like L1 / R2 / Lane L3 -> 'L1', 'R2', 'L3'."""
    for part in parts:
        m = LANE_FOLDER.match(part.strip())
        if m:
            return m.group(1).upper()
    return None


def ocr_available(tesseract_cmd=None):
    """True when the tesseract program can be found (it is installed separately from Python)."""
    try:
        import pytesseract
        if tesseract_cmd:
            pytesseract.pytesseract.tesseract_cmd = tesseract_cmd
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False
