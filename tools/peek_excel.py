"""
peek_excel.py - show the sheets and the first rows of Excel files, local or on S3, to see a new
vendor layout before configuring it in deliverables/dN.py.

    python tools/peek_excel.py "s3://bucket/path/to/file.xlsx"
    python tools/peek_excel.py "s3://bucket/path/to/folder/"          every Excel file below the folder
    python tools/peek_excel.py data/nsv/zone=A/.../d2/file.xlsm --rows 8 --cols 30
"""
import argparse
import io
import sys
from pathlib import Path

import openpyxl

EXCEL = (".xlsx", ".xlsm")


def s3_files(url):
    import boto3
    bucket, _, key = url[len("s3://"):].partition("/")
    s3 = boto3.client("s3")
    if key.lower().endswith(EXCEL):
        return [(url, lambda: s3.get_object(Bucket=bucket, Key=key)["Body"].read())]
    out = []
    for page in s3.get_paginator("list_objects_v2").paginate(Bucket=bucket, Prefix=key):
        for o in page.get("Contents", []):
            k = o["Key"]
            if k.lower().endswith(EXCEL) and "/~$" not in k:
                out.append((f"s3://{bucket}/{k}", lambda k=k: s3.get_object(Bucket=bucket, Key=k)["Body"].read()))
    return out


def local_files(path):
    p = Path(path)
    files = [p] if p.is_file() else sorted(f for f in p.rglob("*") if f.suffix.lower() in EXCEL
                                           and not f.name.startswith("~$"))
    return [(str(f), f.read_bytes) for f in files]


def peek(name, data, rows, cols):
    print(f"\n#################### {name}  ({len(data):,} bytes)")
    if not data.startswith(b"PK"):
        print(f"  NOT an xlsx/xlsm (zip) file - starts with {data[:16]!r}")
        return
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    for ws in wb.worksheets:
        print(f"\n  ===== sheet '{ws.title}'  ({ws.max_row} rows x {ws.max_column} cols)")
        for i, row in enumerate(ws.iter_rows(max_row=rows, values_only=True), 1):
            cells = [f"{openpyxl.utils.get_column_letter(j)}={str(c).strip()[:30]!r}"
                     for j, c in enumerate(row[:cols], 1) if c not in (None, "")]
            print(f"    row {i}: {', '.join(cells) if cells else '(empty)'}")
    wb.close()


def main(args):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("path", help="s3://bucket/key, s3://bucket/folder/, or a local file / folder")
    parser.add_argument("--rows", type=int, default=6, help="rows per sheet (default 6)")
    parser.add_argument("--cols", type=int, default=25, help="columns per row (default 25)")
    o = parser.parse_args(args)
    files = s3_files(o.path) if o.path.startswith("s3://") else local_files(o.path)
    if not files:
        print(f"no .xlsx / .xlsm file under {o.path}")
        return 1
    for name, read in files:
        try:
            peek(name, read(), o.rows, o.cols)
        except Exception as e:
            print(f"\n#################### {name}\n  cannot open: {e}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
